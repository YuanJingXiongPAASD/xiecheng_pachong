"""
携程酒店评论爬虫 - Excel 导出
=============================
将抓取的评论数据导出为格式化的 Excel 文件，包含评论明细和统计两个工作表。
"""

import os
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from config import OUTPUT_DIR, EXCEL_HEADERS, EXCEL_COLUMN_WIDTHS


def export_to_excel(all_reviews: list[dict], hotel_id: str) -> str:
    """
    将评论列表导出为 Excel 文件。

    生成两个工作表：
      - 「酒店评论」：所有评论的明细数据
      - 「统计」：评论总数、平均评分等汇总信息

    Args:
        all_reviews: 评论数据列表
        hotel_id: 酒店 ID，用于文件名

    Returns:
        生成的 Excel 文件路径
    """
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    fp = OUTPUT_DIR / f"酒店评论_{hotel_id}_{ts}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "酒店评论"

    # ---- 样式 ----
    hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    hfont = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ---- 表头 ----
    for ci, h in enumerate(EXCEL_HEADERS, 1):
        c = ws.cell(row=1, column=ci, value=h)
        c.fill = hfill
        c.font = hfont
        c.border = border
        c.alignment = Alignment(horizontal="center", vertical="center")

    # ---- 数据行 ----
    dfont = Font(name="微软雅黑", size=10)
    for ri, r in enumerate(all_reviews, 2):
        vals = [
            ri - 1,
            r.get("reviewId", ""),
            r.get("userName", ""),
            r.get("score", ""),
            r.get("checkInDate", ""),
            r.get("roomType", ""),
            r.get("travelType", ""),
            r.get("content", ""),
            r.get("replyContent", ""),
            r.get("usefulCount", ""),
            r.get("postDate", ""),
        ]
        for ci, v in enumerate(vals, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = dfont
            c.border = border
            c.alignment = Alignment(vertical="top", wrap_text=(ci in (8, 9)))

    # ---- 列宽 ----
    for ci, w in enumerate(EXCEL_COLUMN_WIDTHS, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # ---- 统计页 ----
    ws2 = wb.create_sheet("统计")
    ws2["A1"].font = ws2["B1"].font = Font(bold=True)
    ws2["A1"], ws2["B1"] = "统计项", "数值"

    stats = [
        ("酒店ID", hotel_id),
        ("评论总数", len(all_reviews)),
        ("爬取时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]

    scores = []
    for r in all_reviews:
        try:
            s = float(r.get("score", 0))
            if s > 0:
                scores.append(s)
        except Exception:
            pass
    if scores:
        stats += [
            ("平均评分", f"{sum(scores) / len(scores):.2f}"),
            ("最高/最低", f"{max(scores):.0f} / {min(scores):.0f}"),
        ]

    for ri, (k, v) in enumerate(stats, 2):
        ws2.cell(row=ri, column=1, value=k).font = Font(name="微软雅黑", size=10)
        ws2.cell(row=ri, column=2, value=v).font = Font(name="微软雅黑", size=10)
    ws2.column_dimensions["A"].width = 14
    ws2.column_dimensions["B"].width = 60

    wb.save(fp)

    print(f"\n📊 Excel: {fp}")
    print(f"   共 {len(all_reviews)} 条评论")
    return str(fp)
